#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Génère le cahier de charges de DataVision aux formats PDF et DOCX.

Le contenu est décrit une seule fois (structure CONDUIT ci-dessous), puis
rendu par deux moteurs :
  * reportlab  → Cahier_des_charges_DataVision.pdf
  * python-docx → Cahier_des_charges_DataVision.docx

Usage :  python generer_cahier_des_charges.py
"""

from __future__ import annotations

import datetime as _dt
from pathlib import Path
from xml.sax.saxutils import escape

RACINE = Path(__file__).resolve().parent
FICHIER_PDF = RACINE / "Cahier_des_charges_DataVision.pdf"
FICHIER_DOCX = RACINE / "Cahier_des_charges_DataVision.docx"

# Charte graphique du site (static/css/style.css)
ACCENT = "#4f46e5"
ACCENT_CLAIR = "#f6f7ff"
BORDURE = "#d9dcf0"
TEXTE = "#1f2340"

AUTEUR = "Adem Dridi"
VERSION = "1.0"
DATE = _dt.date.today().strftime("%d %B %Y")


# ======================================================================
#  CONTENU  —  liste de blocs : ("h1"|"h2"|"p"|"ul"|"ol"|"table"|"code"|"page", ...)
# ======================================================================

CONDUIT: list = [
    # ------------------------------------------------------------------
    ("h1", "1. Présentation du projet"),

    ("h2", "1.1 Contexte"),
    ("p", "Les organisations de toutes tailles produisent chaque jour des fichiers "
          "d'exploitation au format CSV ou Excel : ventes, stocks, clients, fournisseurs, "
          "budgets. Ces fichiers sont rarement exploités au-delà de quelques totaux "
          "calculés à la main dans un tableur, faute de temps et de compétences en "
          "programmation ou en statistiques."),
    ("p", "DataVision s'inscrit dans cette dynamique de démocratisation de la donnée "
          "(data self-service) : permettre à une personne qui n'écrit ni SQL ni Python "
          "de transformer un fichier brut en tableau de bord clair, partageable et "
          "exportable, en quelques secondes."),

    ("h2", "1.2 Problématique"),
    ("p", "Le traitement manuel des fichiers de données présente des limites "
          "récurrentes :"),
    ("ul", [
        "Nettoyage long et fastidieux : doublons, dates mal formats, valeurs manquantes, colonnes mal typées.",
        "Analyse réservée aux spécialistes : écrire du code ou des formules complexes pour aller au-delà des totaux.",
        "Résultats non reproductibles : un tableur modifié à la main n'est pas auditable.",
        "Aucune visualisation immédiate : les graphiques doivent être construits un par un.",
        "Difficulté à comparer des périodes ou à suivre une tendance dans le temps.",
        "Rapports figés : produire un document de synthèse demande souvent plusieurs heures.",
    ]),

    ("h2", "1.3 Solution proposée"),
    ("p", "DataVision est une application web qui prend en charge la chaîne complète, "
          "de la donnée brute à la décision :"),
    ("ol", [
        "Import d'un fichier CSV ou Excel (bouton ou glisser-déposer), avec contrôle de taille et de format.",
        "Détection automatique des colonnes (date, produit, catégorie, quantité, prix) et correction manuelle possible.",
        "Nettoyage automatique : doublons, typage des dates, valeurs manquantes, lignes mal formées.",
        "Calcul des indicateurs : chiffre d'affaires, nombre de ventes, panier moyen, produits vendus, catégories.",
        "Tableau de bord interactif : courbe d'évolution, répartition par catégorie, Top 5 / Flop 5, corrélations, anomalies.",
        "Filtres par catégorie et par période, pagination de l'aperçu, comparaison de deux périodes.",
        "Prévision de tendance par régression linéaire (R²), exports CSV/Excel et rapport PDF.",
        "Sauvegarde des analyses par compte, avec recherche, tri et suppression.",
    ]),

    ("h2", "1.4 Périmètre du projet"),
    ("table", "Périmètre du projet", ["Domaine", "Inclus dans la version 1", "Exclus / à venir"], [
        ["Utilisateurs", "Inscription, connexion, réinitialisation du mot de passe, compte de démonstration, administrateur Django",
         "Comptes professionnels, rôles multiples, authentification à deux facteurs"],
        ["Données", "CSV et Excel, 100 Mo et 500 000 lignes maximum, une analyse à la fois",
         "Bases de données distantes, flux temps réel, multi-feuilles Excel"],
        ["Analyse", "KPIs, filtres, tendance, corrélations, anomalies, comparaison de périodes",
         "Modèles de machine learning, prévision multivariée"],
        ["Livrables", "Tableau de bord, exports CSV/Excel, rapport PDF, API JSON",
         "Partage public, alertes e-mail programmées, exports image"],
        ["Infrastructure", "Développement local + déploiement Render (PostgreSQL)",
         "Application mobile, clustering, annuaire LDAP"],
    ]),

    # ------------------------------------------------------------------
    ("h1", "2. Objectifs"),

    ("h2", "2.1 Objectif général"),
    ("p", "Concevoir et développer une plateforme web d'analyse automatique de données, "
          "accessible sans compétence technique, permettant de transformer un fichier "
          "CSV ou Excel en tableau de bord interactif et en rapport exportable."),

    ("h2", "2.2 Objectifs spécifiques"),
    ("ul", [
        "Automatiser l'import, la détection des colonnes et le nettoyage des fichiers.",
        "Calculer fiablement les indicateurs commerciaux clés à partir de données brutes.",
        "Restituer ces indicateurs sous forme de graphiques lisibles et interactifs.",
        "Permettre l'exploration des données : filtres, recherche, pagination, comparaison.",
        "Anticiper l'avenir par une tendance et une prévision expliquée (R²).",
        "Détecter automatiquement les corrélations et les anomalies du jeu de données.",
        "Produire des livrables réutilisables : exports CSV/Excel et rapport PDF.",
        "Sécuriser l'accès : comptes, isolation stricte des données entre utilisateurs.",
        "Assurer la fiabilité sur de gros fichiers (troncature signalée, cache d'analyse).",
        "Offrir une interface responsive, en français, avec thème clair et sombre.",
        "Garantir l'accessibilité : navigation clavier, libellés ARIA, contrastes.",
        "Exposer une API JSON propre, consommable par le front et par une SPA Angular.",
        "Documenter, tester (pytest + tests JS) et déployer le projet de façon reproductible.",
    ]),

    ("h2", "2.3 Indicateurs de succès du projet"),
    ("table", "Indicateurs de succès", ["Indicateur", "Cible"], [
        ["Délai entre l'import et le premier graphique", "moins de 10 secondes pour 50 000 lignes"],
        ["Couverture des tests", "135 tests Python + 70 tests JavaScript, tous au vert"],
        ["Fichiers acceptés", "CSV et Excel jusqu'à 100 Mo / 500 000 lignes"],
        ["Pages responsives", "utilisables de 320 px à 2560 px de large"],
        ["Taux de réussite de l'authentification", "100 % des parcours (inscription, connexion, reset)"],
        ["Disponibilité visée en production", "24 h/24 (hébergement Render)"],
    ]),

    # ------------------------------------------------------------------
    ("h1", "3. Acteurs du système"),
    ("table", "Acteurs du système", ["Acteur", "Rôle", "Droits principaux"], [
        ["Visiteur", "Personne non connectée", "Consulter l'accueil, s'inscrire, se connecter, demander un nouveau mot de passe"],
        ["Utilisateur", "Analyste disposant d'un compte", "Importer des fichiers, consulter et filtrer ses analyses, exporter, générer le rapport PDF, supprimer ses analyses"],
        ["Compte de démonstration", "Compte pré-créé pour la démonstration", "Identiques à l'utilisateur, sur un jeu de données de démonstration"],
        ["Administrateur", "Responsable de la plateforme (Django admin)", "Gérer les comptes et les analyses, consulter l'ensemble des données"],
        ["Moteur d'analyse", "Composant interne (analyzer.py)", "Lire le fichier, nettoyer, calculer les KPIs et les séries renvoyées aux graphiques"],
        ["Fournisseur Google", "Service d'identité externe (à venir)", "Authentifier l'utilisateur via OAuth 2.0"],
    ]),

    # ------------------------------------------------------------------
    ("h1", "4. Parcours utilisateur"),

    ("h2", "4.1 Parcours du visiteur"),
    ("ol", [
        "Arrive sur la page d'accueil et découvre la proposition de valeur.",
        "Clique sur « Créer un compte » (ou « Se connecter »).",
        "Reçoit la confirmation et accède à son espace « Mes analyses ».",
        "Est redirigé vers la page d'import pour charger son premier fichier.",
    ]),

    ("h2", "4.2 Parcours de l'analyste"),
    ("ol", [
        "Importe un fichier CSV/Excel (bouton ou glisser-déposer).",
        "Vérifie, si nécessaire, la correspondance des colonnes détectées automatiquement.",
        "Consulte les KPIs, la courbe d'évolution et la répartition par catégorie.",
        "Filtre par catégorie et par période, puis compare deux périodes.",
        "Explore le classement Top/Flop, les corrélations et les anomalies détectées.",
        "Exporte les données (CSV/Excel) ou génère le rapport PDF.",
        "Retrouve plus tard son analyse dans « Mes analyses » grâce à la recherche et au tri.",
    ]),

    ("h2", "4.3 Parcours de l'administrateur"),
    ("ol", [
        "Se connecte à l'interface d'administration (/admin).",
        "Consulte et gère les comptes utilisateurs et les analyses enregistrées.",
        "Contrôle la suppression d'un compte ou d'une analyse problématique.",
    ]),

    # ------------------------------------------------------------------
    ("h1", "5. Exigences fonctionnelles"),
    ("p", "Chaque exigence est identifiée par un code (FR-xx.n) et une priorité : "
          "Haute (indispensable), Moyenne (importante), Basse (souhaitable)."),

    ("h2", "5.1 FR-01 — Comptes et authentification"),
    ("table", "Exigences FR-01", ["ID", "Exigence", "Priorité"], [
        ["FR-01.1", "Inscription avec nom d'utilisateur, e-mail et mot de passe validé (longueur, similarité, caractère numérique).", "Haute"],
        ["FR-01.2", "Connexion / déconnexion, avec restriction des tentatives échouées (blocage après plusieurs essais).", "Haute"],
        ["FR-01.3", "Case « Se souvenir de moi » qui prolonge réellement la session.", "Moyenne"],
        ["FR-01.4", "Réinitialisation du mot de passe par e-mail (lien à usage unique).", "Haute"],
        ["FR-01.5", "Connexion Google (OAuth 2.0 via django-allauth), activée dès que les clés sont fournies.", "Moyenne"],
        ["FR-01.6", "Compte de démonstration visible ou masquable par variable d'environnement.", "Basse"],
        ["FR-01.7", "Accès à l'administration Django réservé au personnel.", "Haute"],
    ]),

    ("h2", "5.2 FR-02 — Import des fichiers"),
    ("table", "Exigences FR-02", ["ID", "Exigence", "Priorité"], [
        ["FR-02.1", "Import par bouton et par glisser-déposer, formats CSV et Excel acceptés.", "Haute"],
        ["FR-02.2", "Contrôle du type et de la taille (100 Mo maximum) avec message d'erreur explicite.", "Haute"],
        ["FR-02.3", "Lecture robuste : détection d'encodage, séparateurs, lignes mal formées ignorées sans faire échouer l'import.", "Haute"],
        ["FR-02.4", "Lecture rapide (< 1 s pour 50 000 lignes) grâce à un moteur optimisé.", "Moyenne"],
        ["FR-02.5", "Troncature signalée à l'utilisateur au-delà de 500 000 lignes (analyse partielle assumée).", "Moyenne"],
        ["FR-02.6", "Affichage du nom du fichier sélectionné avant envoi.", "Basse"],
    ]),

    ("h2", "5.3 FR-03 — Correspondance des colonnes (mapping)"),
    ("table", "Exigences FR-03", ["ID", "Exigence", "Priorité"], [
        ["FR-03.1", "Détection automatique des rôles : date, produit, catégorie, quantité, prix.", "Haute"],
        ["FR-03.2", "Page de correction manuelle si la détection se trompe (liste déroulante par rôle).", "Haute"],
        ["FR-03.3", "Sauvegarde du mapping sur l'analyse et recalcul des indicateurs.", "Haute"],
        ["FR-03.4", "Message d'aide indiquant les colonnes introuvables et comment les renommer.", "Moyenne"],
    ]),

    ("h2", "5.4 FR-04 — Nettoyage automatique"),
    ("table", "Exigences FR-04", ["ID", "Exigence", "Priorité"], [
        ["FR-04.1", "Conversion des dates (formats courants) et extraction de l'année/mois.", "Haute"],
        ["FR-04.2", "Suppression des doublons et signalement du nombre de lignes retirées.", "Haute"],
        ["FR-04.3", "Gestion des valeurs manquantes et des lignes mal formées, comptabilisées.", "Haute"],
        ["FR-04.4", "Arrondi et typage numérique des montants et des quantités.", "Moyenne"],
        ["FR-04.5", "Affichage d'un message « Données nettoyées automatiquement » récapitulatif.", "Moyenne"],
    ]),

    ("h2", "5.5 FR-05 — Indicateurs (KPIs)"),
    ("table", "Exigences FR-05", ["ID", "Exigence", "Priorité"], [
        ["FR-05.1", "Calcul du chiffre d'affaires total, du nombre de ventes, du panier moyen.", "Haute"],
        ["FR-05.2", "Calcul des produits vendus, du nombre de catégories et des valeurs manquantes.", "Haute"],
        ["FR-05.3", "Recalcul immédiat des indicateurs lorsqu'un filtre est appliqué.", "Haute"],
        ["FR-05.4", "Affichage au format français (espaces de milliers, décimales).", "Moyenne"],
    ]),

    ("h2", "5.6 FR-06 — Tableau de bord et visualisation"),
    ("table", "Exigences FR-06", ["ID", "Exigence", "Priorité"], [
        ["FR-06.1", "Courbe d'évolution du chiffre d'affaires par mois (Chart.js).", "Haute"],
        ["FR-06.2", "Graphique en barres : chiffre d'affaires par catégorie, et camembert de répartition.", "Haute"],
        ["FR-06.3", "Classements Top 5 et Flop 5 des produits.", "Haute"],
        ["FR-06.4", "Repeint automatique des graphiques lors du changement de thème clair/sombre.", "Moyenne"],
        ["FR-06.5", "Graphiques libellés (titre, unité, legend) pour l'accessibilité.", "Moyenne"],
        ["FR-06.6", "Aperçu paginé des données (12 lignes par page) avec recherche instantanée.", "Moyenne"],
    ]),

    ("h2", "5.7 FR-07 — Filtres et exploration"),
    ("table", "Exigences FR-07", ["ID", "Exigence", "Priorité"], [
        ["FR-07.1", "Filtre par catégorie et par plage de dates, cumulable.", "Haute"],
        ["FR-07.2", "Bouton « Réinitialiser » qui restaure l'état complet.", "Moyenne"],
        ["FR-07.3", "Recherche et tri dans la liste des analyses (nom, fichier, date, chiffre d'affaires).", "Moyenne"],
        ["FR-07.4", "Comparaison de deux périodes avec variations en pourcentage (CA, ventes, panier).", "Moyenne"],
        ["FR-07.5", "Avertissement lorsque les deux périodes n'ont pas la même durée.", "Basse"],
    ]),

    ("h2", "5.8 FR-08 — Analyse avancée"),
    ("table", "Exigences FR-08", ["ID", "Exigence", "Priorité"], [
        ["FR-08.1", "Régression linéaire par moindres carrés et affichage de la projection.", "Moyenne"],
        ["FR-08.2", "Indicateur de fiabilité R² avec interprétation lisible (solide / irrégulier).", "Moyenne"],
        ["FR-08.3", "Matrice de corrélation entre variables numériques, triée par force.", "Moyenne"],
        ["FR-08.4", "Détection d'anomalies (écarts anormaux) avec explication en clair.", "Moyenne"],
        ["FR-08.5", "Repli propre : message explicite si l'analyse est impossible (1 seule période, etc.).", "Basse"],
    ]),

    ("h2", "5.9 FR-09 — Rapports et exports"),
    ("table", "Exigences FR-09", ["ID", "Exigence", "Priorité"], [
        ["FR-09.1", "Export CSV des données filtrées.", "Haute"],
        ["FR-09.2", "Export Excel des données filtrées.", "Haute"],
        ["FR-09.3", "Rapport PDF (A4) : en-tête, KPIs, graphiques, conclusions, image de marque.", "Haute"],
        ["FR-09.4", "Exports respectant les filtres actifs.", "Haute"],
        ["FR-09.5", "Boutons d'export inactifs ou explicites tant que l'analyse n'est pas prête.", "Basse"],
    ]),

    ("h2", "5.10 FR-10 — Gestion des analyses"),
    ("table", "Exigences FR-10", ["ID", "Exigence", "Priorité"], [
        ["FR-10.1", "Chaque analyse est rattachée à son propriétaire (isolation stricte).", "Haute"],
        ["FR-10.2", "Liste des analyses avec recherche et tri.", "Haute"],
        ["FR-10.3", "Suppression d'une analyse (et de son fichier) par POST uniquement.", "Haute"],
        ["FR-10.4", "Accès direct par identifiant : renvoi vers la liste si l'analyse n'appartient pas à l'utilisateur.", "Haute"],
        ["FR-10.5", "Messages de confirmation (succès / erreur) visibles par l'utilisateur.", "Moyenne"],
    ]),

    ("h2", "5.11 FR-11 — Interface, thème et accessibilité"),
    ("table", "Exigences FR-11", ["ID", "Exigence", "Priorité"], [
        ["FR-11.1", "Interface entièrement en français.", "Haute"],
        ["FR-11.2", "Thème clair / sombre mémorisé, appliqué avant le rendu (pas de clignotement).", "Moyenne"],
        ["FR-11.3", "Mise en page responsive (desktop, tablette, mobile).", "Haute"],
        ["FR-11.4", "Lien d'évitement, points d'ancrage, libellés ARIA, navigation clavier.", "Moyenne"],
        ["FR-11.5", "Respect de la préférence « mouvement réduit » du système.", "Basse"],
        ["FR-11.6", "Feuille de style dédiée à l'impression.", "Basse"],
    ]),

    ("h2", "5.12 FR-12 — API REST et application Angular"),
    ("table", "Exigences FR-12", ["ID", "Exigence", "Priorité"], [
        ["FR-12.1", "API JSON pour les données du tableau de bord (filtres, pagination).", "Haute"],
        ["FR-12.2", "API JSON pour la comparaison de deux périodes.", "Moyenne"],
        ["FR-12.3", "API JSON listant les analyses du compte (consommée par Angular).", "Moyenne"],
        ["FR-12.4", "Toutes les API protégées par la session et renvoyant 404 en cas d'accès indirect.", "Haute"],
        ["FR-12.5", "Page rendue par Angular hébergée par Django (preuve de concept / SPA).", "Basse"],
        ["FR-12.6", "Point de santé /health pour la surveillance de l'hébergeur.", "Basse"],
    ]),

    ("h2", "5.13 FR-13 — Administration"),
    ("table", "Exigences FR-13", ["ID", "Exigence", "Priorité"], [
        ["FR-13.1", "Administration Django : comptes et analyses consultables, modifiables, supprimables.", "Moyenne"],
        ["FR-13.2", "Audibilité : date de création et propriétaire de chaque analyse.", "Moyenne"],
    ]),

    ("h2", "5.14 Évolutions prévues (hors version 1)"),
    ("ul", [
        "Partage public d'une analyse par lien read-only.",
        "Alertes e-mail programmées (écart au seuil, rapport hebdomadaire).",
        "Import direct depuis une URL ou une connexion à une base de données.",
        "Export des graphiques en PNG/SVG et présentation plein écran.",
        "Comparaison de deux fichiers (et non seulement de deux périodes).",
        "Multi-feuilles Excel et import par onglet.",
        "Traduction de l'interface (anglais/arabe) et formats régionaux.",
        "Version SPA complète d'Angular pour l'ensemble des écrans.",
    ]),

    # ------------------------------------------------------------------
    ("h1", "6. Exigences non fonctionnelles"),
    ("table", "Exigences non fonctionnelles", ["ID", "Catégorie", "Exigence", "Critère mesurable"], [
        ["RNF-01", "Sécurité", "Protection des formulaires (CSRF), échappement systématique des données affichées.",
         "Aucune faille XSS détectée par les tests dédiés"],
        ["RNF-02", "Sécurité", "Mots de passe hachés, validation forte, limitation des tentatives.",
         "Blocage après 5 essais ratés"],
        ["RNF-03", "Sécurité", "Isolation stricte des analyses entre comptes.",
         "Un utilisateur ne peut jamais voir une analyse d'un autre"],
        ["RNF-04", "Sécurité", "Secrets (clés Google, sel de signature) lus dans un fichier .env ignoré par Git.",
         "Aucun secret versionné"],
        ["RNF-05", "Performance", "Temps de réponse du tableau de bord.",
         "Moins de 3 secondes pour 50 000 lignes"],
        ["RNF-06", "Performance", "Mise en cache des analyseurs recalculés.",
         "Rechargement d'une analyse sans relire le fichier"],
        ["RNF-07", "Performance", "Contrôle mémoire : taille d'import et nombre de lignes limités.",
         "100 Mo / 500 000 lignes, troncature signalée"],
        ["RNF-08", "Disponibilité", "Application accessible en continu.",
         "24 h/24 sur l'hébergeur choisi"],
        ["RNF-09", "Compatibilité", "Navigateurs récents (Chrome, Firefox, Edge, Safari).",
         "Fonctions de base disponibles sur les 4"],
        ["RNF-10", "Responsive", "Interface utilisable sur tous les écrans.",
         "De 320 px à 2560 px sans défilement horizontal"],
        ["RNF-11", "Accessibilité", "Navigation clavier, libellés ARIA, contrastes suffisants.",
         "Aucun blocage au test clavier"],
        ["RNF-12", "Maintenabilité", "Code commenté, tests automatisés, intégration continue.",
         "135 tests Python + 70 tests JS au vert à chaque push"],
        ["RNF-13", "Ergonomie", "Interface en français, vocabulaire cohérent, messages d'aide.",
         "Aucun libellé technique non expliqué"],
        ["RNF-14", "Traçabilité", "Chaque analyse conserve sa date et son propriétaire.",
         "Champs obligatoires non modifiables"],
    ]),

    # ------------------------------------------------------------------
    ("h1", "7. Règles de gestion et contraintes"),

    ("h2", "7.1 Règles de gestion"),
    ("table", "Règles de gestion", ["ID", "Règle"], [
        ["RG-01", "Une analyse appartient à un seul utilisateur et n'est visible que par lui."],
        ["RG-02", "Le chiffre d'affaires est la somme de (quantité × prix) des lignes retenues."],
        ["RG-03", "Le panier moyen est le chiffre d'affaires rapporté au nombre de ventes (null si zéro vente)."],
        ["RG-04", "Les filtres (catégorie, période) sont cumulatifs et recalculent tous les indicateurs."],
        ["RG-05", "La comparaison de périodes porte sur des plages de dates strictes, hors bornes."],
        ["RG-06", "La suppression d'une analyse supprime aussi son fichier importé."],
        ["RG-07", "Un fichier non conforme au type ou à la taille est refusé avant tout traitement."],
        ["RG-08", "Le mapping manuel, une fois validé, écrase la détection automatique."],
        ["RG-09", "Au-delà de 500 000 lignes, seules les premières lignes sont analysées et le signalement est affiché."],
        ["RG-10", "Un lien de réinitialisation de mot de passe est valable une seule fois."],
    ]),

    ("h2", "7.2 Contraintes"),
    ("ul", [
        "Formats acceptés : .csv, .xls, .xlsx (encodages UTF-8 et ISO-8859-1 tolérés).",
        "Taille maximale d'un fichier : 100 Mo (configurable par variable d'environnement).",
        "Ligne maximale analysée : 500 000 (configurable).",
        "Base de données SQLite en développement, PostgreSQL en production.",
        "Application web responsive, en français, livrée avec sa documentation.",
    ]),

    # ------------------------------------------------------------------
    ("h1", "8. Product Backlog"),
    ("p", "Le backlog regroupe les user stories priorisées (Haute / Moyenne / Basse) "
          "et leur état d'avancement : Livré, En cours, Prévu."),
    ("table", "Product backlog du projet",
        ["ID", "Fonctionnalité", "User story", "Priorité", "État"], [
        ["1", "Comptes", "En tant qu'utilisateur, je veux créer un compte afin d'accéder à mes analyses.", "Haute", "Livré"],
        ["1.1", "", "En tant qu'utilisateur, je veux me connecter et me déconnecter de manière sécurisée.", "Haute", "Livré"],
        ["1.2", "", "En tant qu'utilisateur, je veux réinitialiser mon mot de passe par e-mail.", "Haute", "Livré"],
        ["1.3", "", "En tant qu'utilisateur, je veux rester connecté grâce à « Se souvenir de moi ».", "Moyenne", "Livré"],
        ["1.4", "", "En tant qu'utilisateur, je veux me connecter avec mon compte Google.", "Moyenne", "En cours"],
        ["1.5", "", "En tant qu'administrateur, je veux gérer les comptes depuis l'interface d'administration.", "Moyenne", "Livré"],
        ["2", "Import", "En tant qu'utilisateur, je veux importer un fichier CSV ou Excel par clic ou glisser-déposer.", "Haute", "Livré"],
        ["2.1", "", "En tant qu'utilisateur, je veux être averti si mon fichier est trop volumineux ou du mauvais format.", "Haute", "Livré"],
        ["2.2", "", "En tant qu'utilisateur, je veux voir le nom du fichier avant l'envoi.", "Basse", "Livré"],
        ["3", "Correspondance des colonnes", "En tant qu'utilisateur, je veux vérifier les colonnes détectées automatiquement.", "Haute", "Livré"],
        ["3.1", "", "En tant qu'utilisateur, je veux corriger manuellement un rôle de colonne.", "Haute", "Livré"],
        ["4", "Nettoyage", "En tant qu'utilisateur, je veux que les doublons et les valeurs manquantes soient traités automatiquement.", "Haute", "Livré"],
        ["4.1", "", "En tant qu'utilisateur, je veux savoir combien de lignes ont été nettoyées.", "Moyenne", "Livré"],
        ["5", "Indicateurs", "En tant qu'utilisateur, je veux voir le chiffre d'affaires, les ventes et le panier moyen.", "Haute", "Livré"],
        ["5.1", "", "En tant qu'utilisateur, je veux que les indicateurs se mettent à jour avec les filtres.", "Haute", "Livré"],
        ["6", "Tableau de bord", "En tant qu'utilisateur, je veux voir l'évolution du chiffre d'affaires dans le temps.", "Haute", "Livré"],
        ["6.1", "", "En tant qu'utilisateur, je veux comparer les catégories entre elles.", "Haute", "Livré"],
        ["6.2", "", "En tant qu'utilisateur, je veux identifier mes produits les plus et les moins vendus.", "Haute", "Livré"],
        ["6.3", "", "En tant qu'utilisateur, je veux consulter un aperçu paginé de mes données.", "Moyenne", "Livré"],
        ["7", "Filtres", "En tant qu'utilisateur, je veux filtrer par catégorie et par période.", "Haute", "Livré"],
        ["7.1", "", "En tant qu'utilisateur, je veux réinitialiser d'un clic tous les filtres.", "Moyenne", "Livré"],
        ["7.2", "", "En tant qu'utilisateur, je veux comparer deux périodes avec les variations en pourcentage.", "Moyenne", "Livré"],
        ["8", "Analyse avancée", "En tant qu'utilisateur, je veux voir la tendance et sa prévision.", "Moyenne", "Livré"],
        ["8.1", "", "En tant qu'utilisateur, je veux comprendre la fiabilité de la prévision (R²).", "Moyenne", "Livré"],
        ["8.2", "", "En tant qu'utilisateur, je veux voir les corrélations entre variables.", "Moyenne", "Livré"],
        ["8.3", "", "En tant qu'utilisateur, je veux être averti des anomalies détectées.", "Moyenne", "Livré"],
        ["9", "Exports", "En tant qu'utilisateur, je veux exporter les données filtrées en CSV.", "Haute", "Livré"],
        ["9.1", "", "En tant qu'utilisateur, je veux exporter les données filtrées en Excel.", "Haute", "Livré"],
        ["9.2", "", "En tant qu'utilisateur, je veux télécharger un rapport PDF complet.", "Haute", "Livré"],
        ["10", "Gestion des analyses", "En tant qu'utilisateur, je veux retrouver toutes mes analyses sauvegardées.", "Haute", "Livré"],
        ["10.1", "", "En tant qu'utilisateur, je veux rechercher et trier mes analyses.", "Moyenne", "Livré"],
        ["10.2", "", "En tant qu'utilisateur, je veux supprimer une analyse et son fichier.", "Haute", "Livré"],
        ["11", "Interface", "En tant qu'utilisateur, je veux basculer entre thème clair et sombre.", "Moyenne", "Livré"],
        ["11.1", "", "En tant qu'utilisateur, je veux utiliser le site sur mobile comme sur PC.", "Haute", "Livré"],
        ["11.2", "", "En tant qu'utilisateur, je veux naviguer au clavier et lire les libellés d'accessibilité.", "Moyenne", "Livré"],
        ["11.3", "", "En tant qu'utilisateur, je veux une interface entièrement en français.", "Haute", "Livré"],
        ["12", "API et Angular", "En tant que développeur, je veux une API JSON pour les données du dashboard.", "Haute", "Livré"],
        ["12.1", "", "En tant que développeur, je veux une API listant mes analyses.", "Moyenne", "Livré"],
        ["12.2", "", "En tant que développeur, je veux une page rendue par Angular hébergée par Django.", "Basse", "Livré"],
        ["12.3", "", "En tant que développeur, je veux une API de comparaison de périodes.", "Moyenne", "Livré"],
        ["13", "Fiabilité", "En tant qu'utilisateur, je veux que les gros fichiers soient traités sans saturer la mémoire.", "Haute", "Livré"],
        ["13.1", "", "En tant qu'utilisateur, je veux que le recalcul d'une analyse soit instantané (cache).", "Moyenne", "Livré"],
        ["13.2", "", "En tant que responsable, je veux une surveillance de l'état du service (/health).", "Basse", "Livré"],
        ["14", "Évolutions", "En tant qu'utilisateur, je veux partager une analyse par lien.", "Moyenne", "Prévu"],
        ["14.1", "", "En tant qu'utilisateur, je veux recevoir une alerte e-mail sur un seuil.", "Basse", "Prévu"],
        ["14.2", "", "En tant qu'utilisateur, je veux exporter un graphique en image.", "Basse", "Prévu"],
    ]),

    # ------------------------------------------------------------------
    ("h1", "9. Sprint Planning"),
    ("table", "Planification des sprints", ["Sprint", "Intitulé", "Fonctionnalités", "Durée"], [
        ["Sprint 1", "Socle et comptes",
         "Projet Django, modèles, inscription, connexion, réinitialisation du mot de passe, isolation des comptes, page d'accueil.", "2 semaines"],
        ["Sprint 2", "Import et nettoyage",
         "Import CSV/Excel, contrôles de taille et de format, détection des colonnes, page de mapping, nettoyage automatique, enregistrement d'une analyse.", "2 semaines"],
        ["Sprint 3", "Tableau de bord et filtres",
         "KPIs, graphiques Chart.js, aperçu paginé, filtres catégorie/période, recherche et tri des analyses, exports CSV/Excel.", "2 semaines"],
        ["Sprint 4", "Analyse avancée et PDF",
         "Tendance et prévision (R²), corrélations, anomalies, comparaison de périodes, rapport PDF.", "2 semaines"],
        ["Sprint 5", "Fiabilité et qualité",
         "Lecteur rapide, troncature, cache des analyseurs, limitation des tentatives, tests (pytest + jsdom), intégration continue.", "2 semaines"],
        ["Sprint 6", "Finitions et Angular",
         "Thème clair/sombre, responsive, accessibilité, icônes, page Angular, documentation, déploiement Render.", "2 semaines"],
    ]),

    # ------------------------------------------------------------------
    ("h1", "10. Architecture technique"),
    ("p", "L'application suit une architecture trois tiers classique, sans étape de "
          "compilation côté serveur : le navigateur reçoit des pages complètes générées "
          "par Django, enrichies en JavaScript pour l'interactivité."),
    ("table", "Architecture par couches", ["Couche", "Composant", "Rôle"], [
        ["Présentation", "Gabarits Django + CSS + JavaScript vanilla", "Affichage des pages, thème, accessibilité, interactions"],
        ["Visualisation", "Chart.js (CDN)", "Courbes, barres, camemberts, repeint selon le thème"],
        ["Icônes", "Bootstrap Icons (CDN)", "Icônes vectorielles de l'interface"],
        ["Interface alternative", "Angular 22 (front/)", "Page /angular rendue côté client et alimentée par les API"],
        ["Web", "Django 6 (vues, formulaires, sessions, admin)", "Routes, contrôle d'accès, rendu, API JSON"],
        ["Métier", "analyzer.py (Pandas / NumPy)", "Lecture, nettoyage, KPIs, séries, régression, corrélations, anomalies"],
        ["Livrable", "reportlab", "Génération du rapport PDF"],
        ["Données", "SQLite (dev) / PostgreSQL (prod)", "Comptes et analyses"],
        ["Cache", "locmem / Redis", "Analyseurs recalculés partagés entre requêtes"],
        ["Exploitation", "gunicorn + Render + GitHub Actions", "Déploiement et intégration continue"],
    ]),
    ("p", "Flux d'un import : choix du fichier → contrôles (type, taille) → détection "
          "des colonnes → nettoyage → calcul des indicateurs → enregistrement de "
          "l'analyse → redirection vers le tableau de bord, qui consomme l'API JSON "
          "pour dessiner les graphiques."),

    # ------------------------------------------------------------------
    ("h1", "11. Modèle de données"),
    ("table", "Entités principales", ["Entité", "Champs principaux", "Remarques"], [
        ["Utilisateur", "identifiant, nom, e-mail, mot de passe haché, date d'inscription, statut administrateur",
         "Table standard de Django, complétée éventuellement par django-allauth"],
        ["Projet (analyse)", "nom, propriétaire, fichier importé, date de création, nombre de lignes, chiffre d'affaires, mapping des colonnes",
         "Une ligne = une analyse sauvegardée ; suppression en cascade avec le compte"],
        ["Session", "clé de session, utilisateur, expiration", "Gérée par Django (cookies signés)"],
        ["Cache d'analyse", "identifiant de projet → analyseur calculé", "Vie courte, purge lors des changements"],
        ["Tentatives de connexion", "adresse IP, utilisateur, compteur, expiration", "Protection contre le force brute"],
    ]),

    # ------------------------------------------------------------------
    ("h1", "12. API REST"),
    ("p", "Toutes les API exigent une session authentifiée et renvoient du JSON."),
    ("table", "API JSON", ["Méthode", "URL", "Rôle", "Réponse"], [
        ["GET", "/api/donnees?id=&categorie=&debut=&fin=&page=", "Données du tableau de bord (KPIs, séries, classements, anomalies)", "JSON complet"],
        ["GET", "/api/comparaison?id=&p1_debut=&p1_fin=&p2_debut=&p2_fin=", "Comparaison de deux périodes avec variations", "JSON"],
        ["GET", "/api/projets", "Analyses du compte connecté (consommé par Angular)", "JSON"],
        ["GET", "/health", "État du service (surveillance de l'hébergeur)", "{ \"statut\": \"ok\" }"],
        ["GET", "/export?format=csv|excel", "Téléchargement des données filtrées", "Fichier"],
        ["GET", "/projet/<id>/rapport", "Rapport PDF", "PDF"],
        ["POST", "/upload", "Import d'un fichier", "Redirection"],
        ["POST", "/projet/<id>/supprimer", "Suppression d'une analyse", "Redirection"],
    ]),

    # ------------------------------------------------------------------
    ("h1", "13. Tableaux de bord"),

    ("h2", "13.1 Tableau de bord de l'utilisateur"),
    ("table", "Contenu du tableau de bord", ["Bloc", "Contenu"], [
        ["Entête", "Nom de l'analyse, fichier, date d'import, boutons Rapport PDF et exports"],
        ["Indicateurs", "Chiffre d'affaires, nombre de ventes, panier moyen, produits vendus, catégories, valeurs manquantes"],
        ["Évolution", "Courbe mensuelle du chiffre d'affaires et prévision de tendance"],
        ["Catégories", "Chiffre d'affaires par catégorie et part relative"],
        ["Classements", "Top 5 et Flop 5 des produits"],
        ["Corrélations", "Paires de variables les plus corrélées, avec force"],
        ["Anomalies", "Écarts inhabituels détectés et expliqués"],
        ["Comparaison", "Deux périodes côte à côte avec variations en pourcentage"],
        ["Données", "Aperçu paginé (12 lignes) avec recherche instantanée"],
    ]),

    ("h2", "13.2 Tableau de bord de l'administrateur"),
    ("table", "Contenu de l'administration", ["Bloc", "Contenu"], [
        ["Comptes", "Liste, création, modification et suppression des utilisateurs"],
        ["Analyses", "Liste de toutes les analyses, propriétaire, date, taille du fichier"],
        ["Journal", "Activité récente, tentatives de connexion bloquées"],
        ["Paramètres", "Variables d'environnement (limites d'upload, identifiants de démonstration)"],
    ]),

    # ------------------------------------------------------------------
    ("h1", "14. Technologies"),
    ("table", "Stack technique", ["Couche", "Technologie", "Usage"], [
        ["Langage", "Python 3", "Logique serveur et moteur d'analyse"],
        ["Framework web", "Django 6", "Routes, templates, formulaires, sessions, admin"],
        ["Analyse", "Pandas + NumPy", "Nettoyage, statistiques, régression"],
        ["Formats", "openpyxl (via pandas)", "Lecture et export Excel"],
        ["Rapports", "reportlab", "Génération du PDF"],
        ["Authentification", "django-allauth + Django auth", "Comptes et connexion Google (OAuth 2.0)"],
        ["Front-end", "JavaScript vanilla", "Filtres, pagination, exports, thème"],
        ["Graphiques", "Chart.js 4", "Visualisations du tableau de bord"],
        ["Icônes", "Bootstrap Icons 1.13", "Icônes de l'interface"],
        ["SPA (preuve de concept)", "Angular 22", "Page /angular et API JSON associées"],
        ["Base de données", "SQLite → PostgreSQL", "Dev → production (dj-database-url)"],
        ["Cache", "locmem → Redis (django-redis)", "Analyseurs mis en cache"],
        ["Serveur", "gunicorn", "Exécution en production"],
        ["Hébergement", "Render", "Application + base managée"],
        ["Tests", "pytest + node:test + jsdom", "135 tests Python, 70 tests JavaScript"],
        ["Intégration continue", "GitHub Actions", "Exécution des tests à chaque push"],
        ["Méthode", "Scrum + Git/GitHub", "Sprints, revue de code, traçabilité"],
    ]),

    # ------------------------------------------------------------------
    ("h1", "15. Méthode et qualité"),
    ("ul", [
        "Méthode agile Scrum : sprints de deux semaines, backlog priorisé, revue en fin de sprint.",
        "Gestion de version sur Git/GitHub, branches par fonctionnalité.",
        "Intégration continue GitHub Actions : les deux suites de tests sont exécutées à chaque push.",
        "Tests unitaires et d'intégration côté serveur (pytest) : moteur, vues, sécurité, gros fichiers, exports.",
        "Tests front-end (node:test + jsdom) : échappement anti-XSS, contrat gabarit/JavaScript, formats, cycle de vie.",
        "Revue systématique des parcours critiques : inscription, connexion, import, export, suppression, réinitialisation.",
        "Documentation vivante : README, cahier de charges, commentaire du code.",
    ]),

    # ------------------------------------------------------------------
    ("h1", "16. Planning et jalons"),
    ("table", "Jalons du projet", ["Jalon", "Livrable", "Durée"], [
        ["J1", "Cadrage, maquettes et mise en place du projet", "Semaine 1"],
        ["J2", "Comptes, sécurité et page d'accueil", "Sprint 1"],
        ["J3", "Import, détection des colonnes et nettoyage", "Sprint 2"],
        ["J4", "Tableau de bord, filtres et exports", "Sprint 3"],
        ["J5", "Analyse avancée et rapport PDF", "Sprint 4"],
        ["J6", "Fiabilité, performance et tests", "Sprint 5"],
        ["J7", "Finitions, Angular, documentation et mise en ligne", "Sprint 6"],
    ]),

    # ------------------------------------------------------------------
    ("h1", "17. Livrables"),
    ("ul", [
        "Code source complet du site (Django + front), versionné sur Git.",
        "Moteur d'analyse indépendant (analyzer.py) documenté.",
        "Application Angular de préproduction (front/) et son build.",
        "Cahier de charges (PDF et DOCX) et présentation du projet.",
        "README : installation, utilisation, tests, déploiement.",
        "Suites de tests automatisées (135 Python, 70 JavaScript) et pipeline CI.",
        "Configuration de déploiement (render.yaml) et variables d'environnement.",
        "Jeu de données de démonstration et compte de démonstration.",
    ]),

    # ------------------------------------------------------------------
    ("h1", "18. Critères d'acceptation"),
    ("table", "Critères d'acceptation", ["Module", "Critère", "Comment vérifier"], [
        ["Comptes", "Inscription, connexion, déconnexion et réinitialisation fonctionnent", "Parcours complet + tests pytest"],
        ["Import", "Un CSV de 50 000 lignes s'importe en moins de 10 s", "Chronométrage sur le jeu de démonstration"],
        ["Mapping", "Une colonne mal détectée se corrige et recalcule les indicateurs", "Modifier un rôle, vérifier les KPIs"],
        ["Nettoyage", "Doublons et valeurs manquantes signalés en clair", "Comparer avant/après dans le message"],
        ["Dashboard", "Les 6 graphiques s'affichent et réagissent aux filtres", "Appliquer un filtre, observer la mise à jour"],
        ["Analyse", "Tendance, corrélations et anomalies affichées ou remplacées par un message explicite", "Jeux de données particuliers"],
        ["Exports", "CSV, Excel et PDF contiennent bien les données filtrées", "Ouvrir les fichiers générés"],
        ["Sécurité", "Aucun accès à une analyse d'un autre compte", "Tentative d'accès direct par identifiant"],
        ["Interface", "Site utilisable sur mobile, thème clair/sombre mémorisé", "Navigateur en mode responsive"],
        ["Qualité", "100 % des tests verts", "python -m pytest -q et npm test"],
        ["Déploiement", "Le site répond en ligne avec /health au vert", "Visite de l'URL de production"],
    ]),

    # ------------------------------------------------------------------
    ("h1", "19. Risques et mesures d'atténuation"),
    ("table", "Risques identifiés", ["Risque", "Impact", "Mesure d'atténuation"], [
        ["Fichier trop volumineux / mémoire saturée", "Élevé", "Limites configurables, troncature signalée, lecteur rapide"],
        ["Détection automatique des colonnes erronée", "Moyen", "Page de correction manuelle sauvegardable"],
        ["Faille XSS via un nom de colonne ou de produit", "Élevé", "Échappement systématique + tests dédiés"],
        ["Compte ou e-mail usurpé", "Élevé", "Validation stricte, mots de passe hachés, limitation des tentatives"],
        ["Pannes d'API externes (Google, CDN)", "Faible", "Boutons conditionnels, le site reste utilisable hors ligne partiel"],
        ["Dérive du périmètre (fonctionnalités non demandées)", "Moyen", "Backlog priorisé, sprints, validation en fin de sprint"],
        ["Non-conformité du rendu entre navigateurs", "Faible", "CSS simple, tests manuels sur 4 navigateurs"],
        ["Perte de données de production", "Élevé", "Sauvegardes de la base (Render) et dépôt Git du code"],
    ]),

    # ------------------------------------------------------------------
    ("h1", "20. Glossaire"),
    ("table", "Glossaire", ["Terme", "Définition"], [
        ["KPI", "Indicateur clé de performance (chiffre d'affaires, panier moyen, etc.)"],
        ["Mapping", "Correspondance attribuée manuellement entre une colonne du fichier et un rôle attendu"],
        ["R²", "Coefficient de détermination d'une régression : part de la variance expliquée (0 à 1)"],
        ["Corrélation", "Mesure de liaison entre deux variables numériques (de -1 à 1)"],
        ["Anomalie", "Valeur ou écart nettement hors du comportement habituel des données"],
        ["CSV", "Fichier de valeurs séparées par des virgules (ou points-virgules)"],
        ["XLSX", "Classeur Excel moderne (format ZIP contenant des feuilles XML)"],
        ["CSRF", "Technique de protection contre les requêtes forgées depuis un autre site"],
        ["OAuth 2.0", "Protocole d'autorisation permettant de se connecter via un tiers (Google)"],
        ["SPA", "Application monopage : le rendu est assuré par le navigateur (Angular)"],
        ["Régession linéaire", "Modèle statistique prévoyant l'évolution d'une variable à partir du temps"],
        ["Sprint", "Itération agile de durée fixe aboutissant à un incrément livrable"],
    ]),

    # ------------------------------------------------------------------
    ("h1", "21. Annexes"),

    ("h2", "Annexe A — Principales URL"),
    ("table", "Routes de l'application", ["URL", "Page"], [
        ["/", "Accueil et import"],
        ["/inscription", "Création de compte"],
        ["/connexion", "Connexion (mot de passe + Google)"],
        ["/deconnexion", "Déconnexion"],
        ["/projets", "Liste des analyses"],
        ["/upload", "Import d'un fichier (POST)"],
        ["/demo", "Chargement de l'analyse de démonstration"],
        ["/dashboard/<id>", "Tableau de bord d'une analyse"],
        ["/projet/<id>/colonnes", "Correspondance des colonnes"],
        ["/projet/<id>/rapport", "Rapport PDF"],
        ["/projet/<id>/supprimer", "Suppression (POST)"],
        ["/mot-de-passe-oublie", "Demande de réinitialisation"],
        ["/angular", "Page rendue par Angular"],
        ["/admin", "Administration Django"],
        ["/health", "État du service"],
    ]),

    ("h2", "Annexe B — Commandes utiles"),
    ("code", [
        "# Lancer le site en développement",
        "python manage.py runserver 127.0.0.1:8765",
        "",
        "# Tests (135 Python + 70 JavaScript)",
        "python -m pytest -q",
        "npm test",
        "",
        "# Recréer le compte de démonstration",
        "python manage.py creer_demo",
        "",
        "# Reconstruire la page Angular",
        "npm run build:angular",
        "",
        "# Migrations de base de données",
        "python manage.py makemigrations && python manage.py migrate",
    ]),

    ("h2", "Annexe C — Comptes et paramètres de test"),
    ("table", "Paramètres de test", ["Élément", "Valeur"], [
        ["Compte de démonstration", "demo / DataVision2026"],
        ["Adresse locale", "http://127.0.0.1:8765"],
        ["Jeu de données", "data/sample_sales.csv (généré par data_generator.py)"],
        ["Clés Google", "GOOGLE_CLIENT_ID et GOOGLE_CLIENT_SECRET dans .env (fichier ignoré par Git)"],
        ["Limite d'import", "DATAVISION_MAX_UPLOAD_MB = 100, DATAVISION_MAX_ROWS = 500000"],
    ]),

    ("h2", "Annexe D — Dépendances principales"),
    ("code", [
        "# Python (requirements.txt)",
        "django, django-allauth, requests, pandas, numpy, openpyxl, reportlab,",
        "dj-database-url, django-redis, gunicorn, pytest",
        "",
        "# JavaScript (package.json)",
        "jsdom  — tests front-end (aucun code produit côté Node)",
        "",
        "# Angular (front/package.json)",
        "@angular/core, @angular/common, rxjs, typescript",
    ]),
]


# ======================================================================
#  RENDU PDF  (reportlab)
# ======================================================================

def rendre_pdf(chemin: Path) -> int:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (BaseDocTemplate, Frame, ListFlowable, ListItem,
                                    PageBreak, PageTemplate, Paragraph, Preformatted,
                                    Spacer, Table, TableStyle)
    from reportlab.platypus.tableofcontents import TableOfContents

    couleur_accent = colors.HexColor(ACCENT)
    couleur_bord = colors.HexColor(BORDURE)
    couleur_fond = colors.HexColor(ACCENT_CLAIR)
    marge = 1.9 * cm

    styles = getSampleStyleSheet()
    s_h1 = ParagraphStyle("Titre1", parent=styles["Heading1"], fontName="Helvetica-Bold",
                          fontSize=15, leading=19, textColor=couleur_accent,
                          spaceBefore=18, spaceAfter=9)
    s_h2 = ParagraphStyle("Titre2", parent=styles["Heading2"], fontName="Helvetica-Bold",
                          fontSize=11.5, leading=15, textColor=colors.HexColor(TEXTE),
                          spaceBefore=13, spaceAfter=6)
    s_corps = ParagraphStyle("Corps", parent=styles["BodyText"], fontName="Helvetica",
                             fontSize=9.7, leading=14, alignment=TA_JUSTIFY,
                             textColor=colors.HexColor(TEXTE), spaceAfter=7)
    s_puce = ParagraphStyle("Puce", parent=s_corps, spaceAfter=3, alignment=TA_JUSTIFY)
    s_legende = ParagraphStyle("Legende", parent=s_corps, fontName="Helvetica-Oblique",
                               fontSize=8.4, textColor=colors.HexColor("#6b7190"),
                               spaceBefore=8, spaceAfter=3, alignment=TA_LEFT)
    s_cell = ParagraphStyle("Cellule", parent=s_corps, fontSize=8.3, leading=11,
                            spaceAfter=0, alignment=TA_LEFT)
    s_cell_t = ParagraphStyle("CelluleEntete", parent=s_cell, fontName="Helvetica-Bold",
                              textColor=colors.white)
    s_code = ParagraphStyle("Code", parent=s_corps, fontName="Courier", fontSize=8.3,
                            leading=11.5, textColor=colors.HexColor(TEXTE))
    s_couverture_titre = ParagraphStyle("CouvertureTitre", parent=styles["Title"],
                                        fontName="Helvetica-Bold", fontSize=34, leading=40,
                                        textColor=couleur_accent, alignment=TA_CENTER,
                                        spaceBefore=60)
    s_couverture_st = ParagraphStyle("CouvertureSous", parent=s_corps, fontSize=14,
                                     leading=20, alignment=TA_CENTER,
                                     textColor=colors.HexColor(TEXTE), spaceBefore=6)

    num_tableau = [0]

    def legende(texte):
        num_tableau[0] += 1
        return Paragraph(f"Tableau {num_tableau[0]} : {escape(texte)}", s_legende)

    def paragraphes_html(texte: str) -> str:
        return escape(texte)

    def construire_tableau(titre, colonnes, lignes, largeurs=None):
        largeur_totale = A4[0] - 2 * marge
        if not largeurs:
            largeurs = [1 / len(colonnes)] * len(colonnes)
        donnees = [[Paragraph(paragraphes_html(c), s_cell_t) for c in colonnes]]
        for ligne in lignes:
            donnees.append([Paragraph(paragraphes_html(str(c)), s_cell) for c in ligne])
        tableau = Table(donnees, colWidths=[largeur_totale * p for p in largeurs],
                        repeatRows=1, hAlign="LEFT")
        tableau.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), couleur_accent),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("GRID", (0, 0), (-1, -1), 0.4, couleur_bord),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, couleur_fond]),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ]))
        return [legende(titre), tableau, Spacer(1, 8)]

    class Cahier(BaseDocTemplate):
        def afterFlowable(self, flux):
            if isinstance(flux, Paragraph):
                style = flux.style.name
                if style == "Titre1":
                    self.notify("TOCEntry", (0, flux.getPlainText(), self.page))
                elif style == "Titre2":
                    self.notify("TOCEntry", (1, flux.getPlainText(), self.page))

    def pied_de_page(canvas, doc):
        canvas.saveState()
        if doc.page > 1:
            canvas.setFont("Helvetica", 7.5)
            canvas.setFillColor(colors.HexColor("#6b7190"))
            canvas.drawString(marge, 1.15 * cm, f"Cahier de charges — DataVision v{VERSION}")
            canvas.drawRightString(A4[0] - marge, 1.15 * cm, f"Page {doc.page}")
            canvas.setStrokeColor(couleur_bord)
            canvas.line(marge, 1.5 * cm, A4[0] - marge, 1.5 * cm)
        canvas.restoreState()

    doc = Cahier(str(chemin), pagesize=A4, leftMargin=marge, rightMargin=marge,
                 topMargin=marge, bottomMargin=2.1 * cm,
                 title=f"Cahier de charges — DataVision", author=AUTEUR,
                 subject="Cahier de charges du projet DataVision")
    cadre = Frame(marge, 2.1 * cm, A4[0] - 2 * marge, A4[1] - marge - 2.1 * cm, id="normal")
    doc.addPageTemplates([PageTemplate(id="principal", frames=[cadre], onPage=pied_de_page)])

    # ---- Couverture ----
    histoire = [
        Spacer(1, 40),
        Paragraph("DataVision", s_couverture_titre),
        Spacer(1, 10),
        Paragraph("Cahier de charges", s_couverture_st),
        Paragraph("Plateforme d'analyse automatique de données CSV et Excel",
                  ParagraphStyle("ss", parent=s_couverture_st, fontSize=11.5, textColor=colors.HexColor("#6b7190"))),
        Spacer(1, 34),
    ]
    meta = [
        ["Projet", "DataVision — application web d'analyse de données"],
        ["Version", VERSION],
        ["Date", DATE],
        ["Auteur", Auteur := AUTEUR],
        ["Technologies", "Django, Pandas, Chart.js, Angular, PostgreSQL"],
        ["Statut", "Document de référence du projet"],
    ]
    t_meta = Table([[Paragraph(f"<b>{escape(c)}</b>", s_cell), Paragraph(escape(d), s_cell)]
                    for c, d in meta], colWidths=[A4[0] * 0.28, A4[0] * 0.5], hAlign="CENTER")
    t_meta.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.4, couleur_bord),
        ("BACKGROUND", (0, 0), (0, -1), couleur_fond),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    histoire += [t_meta, Spacer(1, 60),
                 Paragraph("Document généré automatiquement par <b>generer_cahier_des_charges.py</b>",
                           ParagraphStyle("pied", parent=s_corps, alignment=TA_CENTER,
                                          textColor=colors.HexColor("#6b7190"), fontSize=8.5)),
                 PageBreak()]

    # ---- Table des matières ----
    histoire.append(Paragraph("Table des matières", s_h1))
    toc = TableOfContents()
    toc.levelStyles = [
        ParagraphStyle("TOC1", fontName="Helvetica-Bold", fontSize=9.8, leading=16,
                       textColor=colors.HexColor(TEXTE), leftIndent=0),
        ParagraphStyle("TOC2", fontName="Helvetica", fontSize=9, leading=14,
                       textColor=colors.HexColor("#4a5070"), leftIndent=16),
    ]
    histoire += [toc, PageBreak()]

    # ---- Corps ----
    for bloc in CONDUIT:
        type_bloc = bloc[0]
        if type_bloc == "h1":
            histoire.append(Paragraph(escape(bloc[1]), s_h1))
        elif type_bloc == "h2":
            histoire.append(Paragraph(escape(bloc[1]), s_h2))
        elif type_bloc == "p":
            histoire.append(Paragraph(escape(bloc[1]), s_corps))
        elif type_bloc == "ul":
            histoire.append(ListFlowable(
                [ListItem(Paragraph(escape(i), s_puce), leftIndent=16) for i in bloc[1]],
                bulletType="bullet", start="•", leftIndent=16, bulletFontSize=8,
                bulletOffsetY=0.5, spaceAfter=7))
        elif type_bloc == "ol":
            histoire.append(ListFlowable(
                [ListItem(Paragraph(escape(i), s_puce), leftIndent=18) for i in bloc[1]],
                bulletType="1", leftIndent=18, bulletFontSize=9, spaceAfter=7))
        elif type_bloc == "table":
            _, titre, colonnes, lignes = bloc
            largeurs = None
            if titre.startswith("Product backlog"):
                largeurs = [0.06, 0.16, 0.5, 0.12, 0.16]
            elif len(colonnes) == 4:
                largeurs = [0.1, 0.16, 0.6, 0.14]
            histoire += construire_tableau(titre, colonnes, lignes, largeurs)
        elif type_bloc == "code":
            blocs = "\n".join(bloc[1])
            histoire.append(Spacer(1, 4))
            histoire.append(Preformatted(blocs, s_code))
            histoire.append(Spacer(1, 8))
        elif type_bloc == "page":
            histoire.append(PageBreak())

    doc.multiBuild(histoire)
    try:
        import pymupdf
        return pymupdf.open(str(chemin)).page_count
    except Exception:
        return 0


# ======================================================================
#  RENDU DOCX  (python-docx)
# ======================================================================

def rendre_docx(chemin: Path) -> int:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    def hex_vers_rgb(valeur: str) -> RGBColor:
        v = valeur.lstrip("#")
        return RGBColor(int(v[0:2], 16), int(v[2:4], 16), int(v[4:6], 16))

    accent = hex_vers_rgb(ACCENT)
    texte = hex_vers_rgb(TEXTE)

    def ombrer(cellule, couleur: str):
        element = OxmlElement("w:shd")
        element.set(qn("w:val"), "clear")
        element.set(qn("w:color"), "auto")
        element.set(qn("w:fill"), couleur.lstrip("#"))
        cellule._tc.get_or_add_tcPr().append(element)

    doc = Document()
    base = doc.styles["Normal"]
    base.font.name = "Calibri"
    base.font.size = Pt(10.5)
    base.paragraph_format.space_after = Pt(6)

    for niveau, taille in (("Heading 1", 16), ("Heading 2", 12.5), ("Heading 3", 11)):
        style = doc.styles[niveau]
        style.font.name = "Calibri"
        style.font.size = Pt(taille)
        style.font.bold = True
        style.font.color.rgb = accent
        style.paragraph_format.space_before = Pt(14 if niveau == "Heading 1" else 10)
        style.paragraph_format.space_after = Pt(5)

    # Marge et pied de page avec numéro
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

    def numero_de_page(paragraphe):
        for balise, texte_balise in (("w:fldChar", None), ("w:instrText", "PAGE")):
            element = OxmlElement(balise)
            if balise == "w:fldChar":
                element.set(qn("w:fldCharType"), "begin")
            else:
                element.set(qn("xml:space"), "preserve")
                element.text = " PAGE "
            paragraphe._p.append(element)
        fin = OxmlElement("w:fldChar")
        fin.set(qn("w:fldCharType"), "end")
        paragraphe._p.append(fin)

    pied = section.footer.paragraphs[0]
    pied.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    numero_de_page(pied)

    # ---- Couverture ----
    for _ in range(4):
        doc.add_paragraph()
    titre = doc.add_paragraph()
    titre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = titre.add_run("DataVision")
    run.font.size = Pt(40)
    run.font.bold = True
    run.font.color.rgb = accent

    sous = doc.add_paragraph()
    sous.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sous.add_run("Cahier de charges")
    run.font.size = Pt(18)

    sous2 = doc.add_paragraph()
    sous2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sous2.add_run("Plateforme d'analyse automatique de données CSV et Excel")
    run.font.size = Pt(12)
    run.font.color.rgb = hex_vers_rgb("#6b7190")

    doc.add_paragraph()
    meta = [
        ("Projet", "DataVision — application web d'analyse de données"),
        ("Version", VERSION),
        ("Date", DATE),
        ("Auteur", AUTEUR),
        ("Technologies", "Django, Pandas, Chart.js, Angular, PostgreSQL"),
        ("Statut", "Document de référence du projet"),
    ]
    t_meta = doc.add_table(rows=len(meta), cols=2)
    t_meta.style = "Table Grid"
    for i, (cle, valeur) in enumerate(meta):
        c0, c1 = t_meta.rows[i].cells
        c0.text = cle
        c1.text = valeur
        c0.paragraphs[0].runs[0].font.bold = True
        ombrer(c0, ACCENT_CLAIR)
        for c in (c0, c1):
            for par in c.paragraphs:
                for r in par.runs:
                    r.font.size = Pt(9.5)
    doc.add_page_break()

    # ---- Table des matières (liste synthétique) ----
    doc.add_heading("Table des matières", level=1)
    for bloc in CONDUIT:
        if bloc[0] in ("h1", "h2"):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            if bloc[0] == "h2":
                p.paragraph_format.left_indent = Cm(0.6)
            run = p.add_run(bloc[1])
            run.font.size = Pt(9.5 if bloc[0] == "h2" else 10.5)
            run.font.bold = bloc[0] == "h1"
            run.font.color.rgb = texte
    doc.add_page_break()

    # ---- Corps ----
    numero_tableau = 0
    for bloc in CONDUIT:
        type_bloc = bloc[0]
        if type_bloc == "h1":
            doc.add_heading(bloc[1], level=1)
        elif type_bloc == "h2":
            doc.add_heading(bloc[1], level=2)
        elif type_bloc == "p":
            doc.add_paragraph(bloc[1])
        elif type_bloc == "ul":
            for item in bloc[1]:
                doc.add_paragraph(item, style="List Bullet")
        elif type_bloc == "ol":
            for item in bloc[1]:
                doc.add_paragraph(item, style="List Number")
        elif type_bloc == "table":
            _, titre, colonnes, lignes = bloc
            numero_tableau += 1
            legende = doc.add_paragraph()
            leg_run = legende.add_run(f"Tableau {numero_tableau} : {titre}")
            leg_run.font.size = Pt(8.5)
            leg_run.font.italic = True
            leg_run.font.color.rgb = hex_vers_rgb("#6b7190")
            legende.paragraph_format.space_before = Pt(8)
            legende.paragraph_format.space_after = Pt(3)

            tableau = doc.add_table(rows=1 + len(lignes), cols=len(colonnes))
            tableau.style = "Table Grid"
            for j, entete in enumerate(colonnes):
                cell = tableau.rows[0].cells[j]
                cell.text = str(entete)
                ombrer(cell, ACCENT)
                par = cell.paragraphs[0]
                for r in par.runs:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    r.font.size = Pt(9)
            for i, ligne in enumerate(lignes, start=1):
                for j, valeur in enumerate(ligne):
                    cell = tableau.rows[i].cells[j]
                    cell.text = str(valeur)
                    if i % 2 == 0:
                        ombrer(cell, ACCENT_CLAIR)
                    for par in cell.paragraphs:
                        par.paragraph_format.space_after = Pt(1)
                        for r in par.runs:
                            r.font.size = Pt(9)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
        elif type_bloc == "code":
            for ligne in bloc[1]:
                par = doc.add_paragraph()
                par.paragraph_format.space_after = Pt(0)
                par.paragraph_format.left_indent = Cm(0.4)
                run = par.add_run(ligne if ligne else " ")
                run.font.name = "Courier New"
                run.font.size = Pt(8.5)
                ombrer_table = None
                del ombrer_table
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
        elif type_bloc == "page":
            doc.add_page_break()

    doc.save(str(chemin))
    return len(doc.paragraphs)


# ======================================================================

def main() -> None:
    pages = rendre_pdf(FICHIER_PDF)
    paragraphes = rendre_docx(FICHIER_DOCX)
    print(f"PDF  : {FICHIER_PDF}  ({pages} pages)")
    print(f"DOCX : {FICHIER_DOCX}  ({paragraphes} paragraphes)")
    print(f"Sections : {sum(1 for b in CONDUIT if b[0] == 'h1')}, "
          f"tableaux : {sum(1 for b in CONDUIT if b[0] == 'table')}")


if __name__ == "__main__":
    main()
